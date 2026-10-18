import io
import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from functools import cache
from pathlib import Path
from typing import Any

from PIL import Image as PILImage
from PIL.Image import Image

from app.db.model import Modality, SourceType

JSON_LEAF_CHARS = 2000
PAGE_DPI        = 110
HEADING_RE      = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
MAC_SOFFICE     = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

Headings = list[tuple[int, str]]
# x0, y0, x1, y1 as fractions of the page, origin top-left
Box = tuple[float, float, float, float]


@dataclass
class Element:
    modality     : Modality
    content      : str
    page         : int | None       = None
    heading_path : list[str] | None = None
    json_path    : str | None       = None
    image        : Image | None     = None
    boxes        : list[Box] | None = None


def parse(data: bytes, filename: str, source_type: SourceType, parser: str) -> list[Element]:
    elements = _parse(data, filename, source_type, parser)

    # Postgres text can't hold NUL, which some PDF text layers contain
    for element in elements:
        element.content = element.content.replace("\x00", "")
        if element.heading_path:
            element.heading_path = [h.replace("\x00", "") for h in element.heading_path]

    return elements


def _parse(data: bytes, filename: str, source_type: SourceType, parser: str) -> list[Element]:
    if source_type in (SourceType.PDF, SourceType.PPTX) and parser in ("docling", "vlm"):
        return parse_docling(data, filename, keep_images=parser == "vlm")

    if source_type in (SourceType.PDF, SourceType.PPTX) and parser == "page":
        return parse_pages(data if source_type == SourceType.PDF else _pptx_to_pdf(data))

    match source_type:
        case SourceType.PDF:
            return parse_pdf(data)
        case SourceType.PPTX:
            return parse_pptx(data)
        case SourceType.MARKDOWN:
            return parse_markdown(_decode(data))
        case SourceType.JSON:
            return parse_json(_decode(data))
        case SourceType.TEXT:
            return [Element(Modality.TEXT, _decode(data))]

    raise ValueError(f"Unsupported source type: {source_type}")


def parse_pdf(data: bytes) -> list[Element]:
    import pymupdf

    elements = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for number, page in enumerate(pdf, start=1):
            width, height = page.rect.width, page.rect.height
            for x0, y0, x1, y1, text, _, kind in page.get_text("blocks"):
                if kind == 0 and text.strip():
                    elements.append(Element(Modality.TEXT, text.strip(), page=number, boxes=[_box(x0, y0, x1, y1, width, height)]))

    return elements


def parse_pages(pdf: bytes) -> list[Element]:
    """One image chunk per page, with the page's text layer as content for full-text search and reranking."""
    import pymupdf

    elements = []
    with pymupdf.open(stream=pdf, filetype="pdf") as document:
        for number, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(dpi=PAGE_DPI)
            image  = PILImage.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            elements.append(Element(Modality.IMAGE, page.get_text().strip(), page=number, image=image, boxes=[(0.0, 0.0, 1.0, 1.0)]))

    return elements


def _pptx_to_pdf(data: bytes) -> bytes:
    """python-pptx can't draw slides, so LibreOffice renders them to PDF first."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice") or (MAC_SOFFICE if Path(MAC_SOFFICE).exists() else None)
    if soffice is None:
        raise RuntimeError("The page parser needs LibreOffice for PPTX, run: brew install --cask libreoffice")

    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / "slides.pptx"
        source.write_bytes(data)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(source)], check=True, capture_output=True, timeout=300)
        return (Path(tmp) / "slides.pdf").read_bytes()


def parse_pptx(data: bytes) -> list[Element]:
    from pptx import Presentation

    elements = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        for shape in slide.shapes:
            if shape.has_table:
                texts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)

        if text := "\n".join(t for t in texts if t.strip()).strip():
            elements.append(Element(Modality.TEXT, text, page=number))

    return elements


def parse_markdown(text: str) -> list[Element]:
    """One element per section, with the path of headings above it."""
    elements : list[Element] = []
    headings : Headings      = []
    lines    : list[str]     = []
    in_code  = False

    def flush() -> None:
        if body := "\n".join(lines).strip():
            elements.append(Element(Modality.TEXT, body, heading_path=_path(headings)))
        lines.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code

        match = None if in_code else HEADING_RE.match(line)
        if match:
            flush()
            _push(headings, len(match.group(1)), match.group(2))
        lines.append(line)

    flush()
    return elements


def parse_json(text: str) -> list[Element]:
    elements: list[Element] = []
    _walk_json(json.loads(text), "$", elements)
    return elements


def _walk_json(value: Any, path: str, elements: list[Element]) -> None:
    """Split into the largest subtrees that fit in JSON_LEAF_CHARS."""
    serialized = json.dumps(value, ensure_ascii=False, indent=1)

    if len(serialized) <= JSON_LEAF_CHARS or not isinstance(value, (dict, list)) or not value:
        elements.append(Element(Modality.TEXT, serialized, json_path=path))
        return

    children = value.items() if isinstance(value, dict) else enumerate(value)
    for key, child in children:
        child_path = f"{path}[{key}]" if isinstance(key, int) else f"{path}.{key}"
        _walk_json(child, child_path, elements)


def parse_docling(data: bytes, filename: str, keep_images: bool) -> list[Element]:
    from docling.datamodel.document  import DocumentStream
    from docling_core.types.doc      import PictureItem, SectionHeaderItem, TableItem, TextItem, TitleItem

    result   = _docling_converter().convert(DocumentStream(name=filename, stream=io.BytesIO(data)))
    document = result.document

    elements : list[Element] = []
    headings : Headings      = []

    for item, _ in document.iterate_items():
        page, boxes = _docling_location(item, document)

        if isinstance(item, (TitleItem, SectionHeaderItem)):
            _push(headings, 0 if isinstance(item, TitleItem) else item.level, item.text)
            elements.append(Element(Modality.TEXT, item.text, page, _path(headings), boxes=boxes))
        elif isinstance(item, TableItem):
            elements.append(Element(Modality.TABLE, item.export_to_markdown(doc=document), page, _path(headings), boxes=boxes))
        elif isinstance(item, PictureItem):
            if keep_images and (image := item.get_image(document)):
                elements.append(Element(Modality.IMAGE, item.caption_text(document), page, _path(headings), image=image, boxes=boxes))
        elif isinstance(item, TextItem) and item.text.strip():
            elements.append(Element(Modality.TEXT, item.text, page, _path(headings), boxes=boxes))

    return elements


def _docling_location(item, document) -> tuple[int | None, list[Box] | None]:
    """Page of the item's first box, and its boxes on that page."""
    if not getattr(item, "prov", None):
        return None, None

    page = item.prov[0].page_no
    if page not in document.pages:
        return page, None

    size  = document.pages[page].size
    boxes = []
    for prov in item.prov:
        if prov.page_no == page:
            b = prov.bbox.to_top_left_origin(size.height)
            boxes.append(_box(b.l, b.t, b.r, b.b, size.width, size.height))

    return page, boxes


def _box(x0: float, y0: float, x1: float, y1: float, width: float, height: float) -> Box:
    clamp = lambda v: round(min(max(v, 0.0), 1.0), 4)
    return clamp(x0 / width), clamp(y0 / height), clamp(x1 / width), clamp(y1 / height)


def _push(headings: Headings, level: int, text: str) -> None:
    while headings and headings[-1][0] >= level:
        headings.pop()
    headings.append((level, text.strip()))


def _path(headings: Headings) -> list[str] | None:
    return [text for _, text in headings] or None


@cache
def _docling_converter():
    from docling.datamodel.base_models      import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter         import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions(generate_picture_images=True, images_scale=2.0)
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})


def _decode(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")
